"""
Word文档生成器
将分类后的文献内容生成格式化的Word文档
"""

import os
from datetime import datetime
from typing import Dict, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


class WordGenerator:
    """Word文档生成器"""

    def __init__(self):
        """初始化Word生成器"""
        self.doc = None

    def generate(
        self,
        disease: str,
        classified_content: Dict[str, List[Dict]],
        module_summaries: Dict[str, str],
        output_path: str = None
    ) -> str:
        """
        生成Word文档

        Args:
            disease: 疾病名称
            classified_content: 分类后的文献内容
            module_summaries: 各模块的综合摘要
            output_path: 输出文件路径

        Returns:
            生成的文档路径
        """
        print("\n开始生成Word文档...")

        # 创建文档
        self.doc = Document()

        # 设置文档样式
        self._setup_document_styles()

        # 添加标题
        self._add_title(disease)

        # 添加文档信息
        self._add_document_info(classified_content)

        # 添加各模块内容
        from .content_classifier import ContentClassifier
        modules = ContentClassifier.MODULES

        for module_key, module_info in modules.items():
            module_name = module_info['name']
            module_desc = module_info['description']
            content = classified_content.get(module_key, [])
            summary = module_summaries.get(module_key, '')

            self._add_module_section(
                module_name,
                module_desc,
                content,
                summary
            )

        # 添加参考文献列表
        self._add_references(classified_content)

        # 保存文档
        if not output_path:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = f"output/{disease}_文献汇总_{timestamp}.docx"

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        self.doc.save(output_path)

        print(f"文档已生成: {output_path}")
        return output_path

    def _setup_document_styles(self):
        """设置文档样式"""
        # 设置中文字体
        self.doc.styles['Normal'].font.name = '宋体'
        self.doc.styles['Normal']._element.rPr.rFonts.set(
            qn('w:eastAsia'), '宋体'
        )
        self.doc.styles['Normal'].font.size = Pt(10.5)

    def _add_title(self, disease: str):
        """添加文档标题"""
        title = self.doc.add_heading(f'{disease}文献检索与汇总报告', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 添加副标题
        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(f'生成日期: {datetime.now().strftime("%Y年%m月%d日")}')
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(100, 100, 100)

        self.doc.add_paragraph()  # 空行

    def _add_document_info(self, classified_content: Dict):
        """添加文档概览信息"""
        self.doc.add_heading('文献概览', level=1)

        # 统计信息
        total_articles = sum(len(content) for content in classified_content.values())
        unique_articles = len(set(
            item['article']['pmid'] if 'pmid' in item['article'] else item['article']['url']
            for content in classified_content.values()
            for item in content
        ))

        info = self.doc.add_paragraph()
        info.add_run(f'检索文献总数: {unique_articles} 篇\n').bold = True

        # 各模块文献数量
        from .content_classifier import ContentClassifier
        modules = ContentClassifier.MODULES

        table = self.doc.add_table(rows=1, cols=2)
        table.style = 'Light Grid Accent 1'

        # 表头
        header_cells = table.rows[0].cells
        header_cells[0].text = '模块'
        header_cells[1].text = '文献数量'

        # 数据行
        for module_key, module_info in modules.items():
            row_cells = table.add_row().cells
            row_cells[0].text = module_info['name']
            row_cells[1].text = str(len(classified_content.get(module_key, [])))

        self.doc.add_paragraph()  # 空行

    def _add_module_section(
        self,
        module_name: str,
        module_desc: str,
        content: List[Dict],
        summary: str
    ):
        """添加模块内容"""
        # 模块标题
        self.doc.add_heading(f'{module_name}', level=1)

        # 模块描述
        desc_para = self.doc.add_paragraph()
        desc_run = desc_para.add_run(f'[{module_desc}]')
        desc_run.font.italic = True
        desc_run.font.color.rgb = RGBColor(100, 100, 100)

        # 综合摘要
        if summary:
            self.doc.add_heading('综合摘要', level=2)
            summary_para = self.doc.add_paragraph(summary)
            summary_para.paragraph_format.first_line_indent = Pt(21)  # 首行缩进

        # 详细文献列表
        if content:
            self.doc.add_heading('相关文献', level=2)

            for idx, item in enumerate(content[:20], 1):  # 最多显示20篇
                article = item['article']
                article_summary = item.get('summary', '')
                key_points = item.get('key_points', [])

                # 文献序号和标题
                title_para = self.doc.add_paragraph()
                title_run = title_para.add_run(f'{idx}. {article.get("title", "无标题")}')
                title_run.bold = True
                title_run.font.size = Pt(11)

                # 文献信息
                info_para = self.doc.add_paragraph()
                info_text = []

                if article.get('authors'):
                    authors = ', '.join(article['authors'][:3])
                    if len(article['authors']) > 3:
                        authors += ' 等'
                    info_text.append(f'作者: {authors}')

                if article.get('journal'):
                    info_text.append(f'期刊: {article["journal"]}')

                if article.get('pub_date'):
                    info_text.append(f'日期: {article["pub_date"]}')

                info_para.add_run(' | '.join(info_text))
                info_para.paragraph_format.left_indent = Pt(21)

                # 文献URL
                if article.get('url'):
                    url_para = self.doc.add_paragraph()
                    url_para.add_run(f'链接: {article["url"]}')
                    url_para.paragraph_format.left_indent = Pt(21)

                # 相关性摘要
                if article_summary:
                    summary_para = self.doc.add_paragraph()
                    summary_para.add_run('摘要: ').bold = True
                    summary_para.add_run(article_summary)
                    summary_para.paragraph_format.left_indent = Pt(21)

                # 关键点
                if key_points:
                    key_para = self.doc.add_paragraph()
                    key_para.add_run('关键点:').bold = True
                    key_para.paragraph_format.left_indent = Pt(21)

                    for point in key_points[:5]:  # 最多5个关键点
                        point_para = self.doc.add_paragraph(
                            f'• {point}',
                            style='List Bullet'
                        )
                        point_para.paragraph_format.left_indent = Pt(42)

                self.doc.add_paragraph()  # 文献之间空行

        else:
            no_content = self.doc.add_paragraph('暂无相关文献。')
            no_content.paragraph_format.left_indent = Pt(21)

        self.doc.add_paragraph()  # 模块之间空行

    def _add_references(self, classified_content: Dict):
        """添加所有参考文献列表"""
        self.doc.add_heading('完整参考文献列表', level=1)

        # 收集所有唯一的文献
        all_articles = {}
        for content in classified_content.values():
            for item in content:
                article = item['article']
                article_id = article.get('pmid') or article.get('url')
                if article_id not in all_articles:
                    all_articles[article_id] = article

        # 按来源分组
        pubmed_articles = []
        cma_articles = []

        for article in all_articles.values():
            if article.get('source') == 'PubMed':
                pubmed_articles.append(article)
            else:
                cma_articles.append(article)

        # PubMed文献
        if pubmed_articles:
            self.doc.add_heading('PubMed文献', level=2)
            for idx, article in enumerate(pubmed_articles, 1):
                self._add_reference_item(idx, article)

        # 中华医学会文献
        if cma_articles:
            self.doc.add_heading('中华医学会文献', level=2)
            for idx, article in enumerate(cma_articles, 1):
                self._add_reference_item(idx, article)

    def _add_reference_item(self, idx: int, article: Dict):
        """添加单个参考文献条目"""
        ref_para = self.doc.add_paragraph()

        # 序号
        ref_para.add_run(f'[{idx}] ').bold = True

        # 作者
        if article.get('authors'):
            authors = ', '.join(article['authors'][:6])
            if len(article['authors']) > 6:
                authors += ', et al'
            ref_para.add_run(f'{authors}. ')

        # 标题
        ref_para.add_run(f'{article.get("title", "无标题")}. ')

        # 期刊和日期
        if article.get('journal'):
            ref_para.add_run(f'{article["journal"]}. ')
        if article.get('pub_date'):
            ref_para.add_run(f'{article["pub_date"]}. ')

        # URL
        if article.get('url'):
            ref_para.add_run(f'{article["url"]}')

        ref_para.paragraph_format.left_indent = Pt(21)
        ref_para.paragraph_format.first_line_indent = Pt(-21)  # 悬挂缩进
